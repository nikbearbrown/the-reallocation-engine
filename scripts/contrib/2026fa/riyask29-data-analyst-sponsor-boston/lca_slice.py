#!/usr/bin/env python3
"""lca_slice.py — cut a DOL LCA disclosure file down to employer x SOC counts for one state.

Closes recipe TODO 1: sponsorship evidence by occupation code instead of by title keyword.
The DOL Office of Foreign Labor Certification publishes every Labor Condition Application
(the filing an employer must make before an H-1B petition) with its SOC code. This script
reads one quarterly disclosure workbook and writes, per employer and SOC code, how many
certified H-1B LCAs had a worksite in the chosen state(s).

It never writes a case number or a person (contact, agent or attorney names, emails, phones):
only employer names, SOC codes, job-title counts, wages and dates. The raw workbook stays
outside the repository (it is ~100+ MB and not gitignored here).

    python3 lca_slice.py --xlsx ~/Downloads/LCA_Disclosure_Data_FY2026_Q3.xlsx \
        --state MA --out data/lca-ma-analyst-fy2026q3.csv --audit data/lca-ma-analyst-fy2026q3-audit.md

Conformance halts (exit 2): a required column missing, or zero rows read. Everything else is
counted in the audit for a person to read. Needs openpyxl.
"""
import argparse
import csv
import json
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
REQUIRED = ["CASE_STATUS", "VISA_CLASS", "JOB_TITLE", "SOC_CODE", "SOC_TITLE", "EMPLOYER_NAME",
            "WORKSITE_CITY", "WORKSITE_STATE", "WAGE_RATE_OF_PAY_FROM", "WAGE_UNIT_OF_PAY", "DECISION_DATE"]
ANNUAL = {"year": 1, "hour": 2080, "week": 52, "bi-weekly": 26, "month": 12}


def annual(rate, unit):
    try:
        r = float(str(rate).replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None
    k = ANNUAL.get(str(unit or "").strip().lower())
    return round(r * k) if k and r > 0 else None


def soc6(code):
    """'15-2051.01' -> '15-2051.01'; '15-2051' -> '15-2051.00'; anything else unchanged."""
    c = str(code or "").strip()
    return c + ".00" if len(c) == 7 and c[2] == "-" else c


def iso(d):
    if isinstance(d, (datetime, date)):
        return d.strftime("%Y-%m-%d")
    return str(d or "")[:10]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--state", nargs="+", default=["MA"])
    ap.add_argument("--families", type=Path, default=HERE / "config" / "lca-soc-families.json")
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--audit", required=True, type=Path)
    a = ap.parse_args()

    from openpyxl import load_workbook          # imported here so --help works without it

    fam = json.loads(a.families.read_text())
    core = {soc6(c) for c in fam["core"]}
    adjacent = {soc6(c) for c in fam["adjacent"]}
    wanted = core | adjacent
    states = {s.upper() for s in a.state}

    wb = load_workbook(a.xlsx, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().upper() for h in next(rows)]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        print(f"[lca_slice] HALT: required column(s) missing: {missing} (record layout changed?)", file=sys.stderr)
        sys.exit(2)
    ix = {c: header.index(c) for c in REQUIRED}

    n = Counter()
    groups = defaultdict(lambda: {"titles": Counter(), "cities": Counter(), "wages": [], "dates": [], "soc_title": ""})
    status_seen, visa_seen = Counter(), Counter()
    for r in rows:
        if all(v is None or str(v).strip() == "" for v in r):
            n["blank"] += 1          # formatted-but-empty rows at the sheet's end; not filings
            continue
        n["read"] += 1
        g = lambda c: r[ix[c]]
        status, visa = str(g("CASE_STATUS") or "").strip(), str(g("VISA_CLASS") or "").strip()
        status_seen[status] += 1
        visa_seen[visa] += 1
        if visa != "H-1B":
            n["drop: not H-1B"] += 1
            continue
        if str(g("WORKSITE_STATE") or "").strip().upper() not in states:
            n["drop: worksite not in " + "/".join(sorted(states))] += 1
            continue
        if status != "Certified":
            n[f"drop: status {status or '(blank)'}"] += 1
            continue
        code = soc6(g("SOC_CODE"))
        if code not in wanted:
            n["drop: SOC not in analyst families"] += 1
            continue
        n["kept"] += 1
        key = (str(g("EMPLOYER_NAME") or "").strip(), code)
        grp = groups[key]
        grp["titles"][str(g("JOB_TITLE") or "").strip().title()] += 1
        grp["cities"][str(g("WORKSITE_CITY") or "").strip().title()] += 1
        grp["soc_title"] = str(g("SOC_TITLE") or "").strip()
        w = annual(g("WAGE_RATE_OF_PAY_FROM"), g("WAGE_UNIT_OF_PAY"))
        if w:
            grp["wages"].append(w)
        grp["dates"].append(iso(g("DECISION_DATE")))
    if not n["read"]:
        print("[lca_slice] HALT: zero rows read", file=sys.stderr)
        sys.exit(2)

    out = []
    for (emp, code), grp in groups.items():
        ds = sorted(d for d in grp["dates"] if d)
        out.append({"employer_name": emp, "soc_code": code, "soc_title": grp["soc_title"],
                    "family": "core" if code in core else "adjacent",
                    "certified_lcas": sum(grp["titles"].values()),
                    "top_job_titles": "; ".join(t for t, _ in grp["titles"].most_common(3)),
                    "worksite_cities": "; ".join(c for c, _ in grp["cities"].most_common(2)),
                    "median_annual_wage_offered": int(statistics.median(grp["wages"])) if grp["wages"] else "",
                    "first_decision": ds[0] if ds else "", "last_decision": ds[-1] if ds else "",
                    "source_file": a.xlsx.name})
    out.sort(key=lambda d: (-d["certified_lcas"], d["employer_name"], d["soc_code"]))
    a.out.parent.mkdir(parents=True, exist_ok=True)
    with a.out.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()) if out else ["employer_name"])
        w.writeheader()
        w.writerows(out)

    emps = Counter()
    for d in out:
        emps[d["employer_name"]] += d["certified_lcas"]
    by_soc = Counter()
    for d in out:
        by_soc[(d["soc_code"], d["soc_title"], d["family"])] += d["certified_lcas"]
    L = ["# LCA slice audit", "", "## Executive summary", "",
         f"This file counts what happened when one Department of Labor H-1B filing file was cut down to "
         f"certified filings for analyst-type occupations at worksites in {', '.join(sorted(states))}. "
         f"It read {n['read']:,} filings and kept {n['kept']:,}, from {len(emps):,} employers. "
         "It is a report for a person to read, not a pass/fail check: a high drop count is expected, "
         "because most filings are for other occupations and states.", "",
         "## Run record", "",
         f"- Source: `{a.xlsx.name}` (DOL OFLC LCA disclosure data), read {date.today().isoformat()}",
         f"- Families: `{a.families.name}` — core {sorted(core)}, adjacent {sorted(adjacent)}",
         f"- Output: `{a.out.name}` ({len(out)} employer × SOC rows; no case numbers, no person's name or contact)", "",
         "## Rows in, rows out", "", "| step | rows |", "|---|---|",
         f"| blank sheet rows skipped (not filings) | {n['blank']:,} |", f"| read | {n['read']:,} |"]
    L += [f"| {k} | {v:,} |" for k, v in sorted(n.items()) if k.startswith("drop")]
    L += [f"| **kept** | **{n['kept']:,}** |", "", "## Kept, by SOC code", "", "| SOC | title | family | certified LCAs |", "|---|---|---|---|"]
    L += [f"| {c} | {t} | {fam_} | {v:,} |" for (c, t, fam_), v in by_soc.most_common()]
    L += ["", "## Top 15 employers (all kept SOCs)", "", "| employer | certified LCAs |", "|---|---|"]
    L += [f"| {e} | {v:,} |" for e, v in emps.most_common(15)]
    L += ["", "## Values seen (for a person to sanity-check the filters)", "",
          "- CASE_STATUS: " + ", ".join(f"{k or '(blank)'} {v:,}" for k, v in status_seen.most_common()),
          "- VISA_CLASS: " + ", ".join(f"{k or '(blank)'} {v:,}" for k, v in visa_seen.most_common()), "",
          "## What this cannot tell you", "",
          "- An LCA is filed before a petition; it is not an approved H-1B, and one LCA can cover several workers.",
          "- A count of zero means no certified analyst-family LCA under this exact employer name in this file and state, "
          "not that the employer never sponsors.",
          "- Wages are the offered rate on the filing, annualised; not what any one person was paid.", ""]
    a.audit.write_text("\n".join(L))
    print(f"[lca_slice] blank rows skipped {n['blank']:,} · read {n['read']:,} · kept {n['kept']:,} · {len(emps):,} employers · {len(out)} rows → {a.out}")


if __name__ == "__main__":
    main()
